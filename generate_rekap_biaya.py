"""Generate Rekap Biaya Bulanan - Web Edukasi Catin Puskesmas Cakung"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title
title = doc.add_heading('REKAP BIAYA BULANAN', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Web Edukasi Calon Pengantin — Puskesmas Cakung')
run.bold = True
run.font.size = Pt(14)

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run('Periode: Per bulan | URL: https://iamwfs30.github.io/educatin-puskes-cakung/').font.size = Pt(10)

doc.add_paragraph()

# === Section A: Infrastruktur ===
doc.add_heading('A. Biaya Infrastruktur', level=1)

table_a = doc.add_table(rows=8, cols=5)
table_a.style = 'Table Grid'
table_a.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_a = ['No', 'Komponen', 'Layanan', 'Keterangan', 'Biaya/bulan']
for i, h in enumerate(headers_a):
    cell = table_a.rows[0].cells[i]
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True

data_a = [
    ['1', 'Hosting & CDN', 'GitHub Pages', 'Static hosting, 100GB bandwidth/bulan', 'Rp 0'],
    ['2', 'Database Rekap', 'Google Sheets + Apps Script', 'Penyimpanan data pasien & skor', 'Rp 0'],
    ['3', 'Domain', 'Subdomain GitHub', 'iamwfs30.github.io', 'Rp 0'],
    ['4', 'SSL/HTTPS', 'GitHub Pages', 'Sertifikat SSL otomatis', 'Rp 0'],
    ['5', 'Icon Library', 'Lucide Icons', 'Open-source, via CDN', 'Rp 0'],
    ['6', 'Notifikasi WhatsApp', 'wa.me (Click-to-Chat)', 'Kirim hasil ke WA admin', 'Rp 0'],
    ['', '', '', 'Subtotal Infrastruktur', 'Rp 0'],
]

for row_idx, row_data in enumerate(data_a):
    for col_idx, val in enumerate(row_data):
        table_a.rows[row_idx + 1].cells[col_idx].text = val

# Bold subtotal row
for cell in table_a.rows[7].cells:
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True

doc.add_paragraph()

# === Section B: Maintenance & Support ===
doc.add_heading('B. Biaya Maintenance & Support', level=1)

table_b = doc.add_table(rows=7, cols=5)
table_b.style = 'Table Grid'
table_b.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_b = ['No', 'Item', 'Deskripsi', 'Frekuensi', 'Biaya/bulan']
for i, h in enumerate(headers_b):
    cell = table_b.rows[0].cells[i]
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True

data_b = [
    ['1', 'Bug Fix', 'Perbaikan error/bug yang ditemukan saat operasional', 'Maks 4 incident/bulan', 'Rp 300.000'],
    ['2', 'Update Konten', 'Ganti/tambah soal, update video, ubah teks', 'Maks 2x/bulan', 'Rp 200.000'],
    ['3', 'Monitoring', 'Cek ketersediaan website, pastikan berjalan normal', 'Mingguan', 'Rp 150.000'],
    ['4', 'Backup Data', 'Export rekap Google Sheets, backup source code', 'Bulanan', 'Rp 100.000'],
    ['5', 'Technical Support', 'Bantuan teknis via WhatsApp (jam kerja)', 'Unlimited chat', 'Rp 250.000'],
    ['', '', '', 'Subtotal Maintenance & Support', 'Rp 1.000.000'],
]

for row_idx, row_data in enumerate(data_b):
    for col_idx, val in enumerate(row_data):
        table_b.rows[row_idx + 1].cells[col_idx].text = val

for cell in table_b.rows[6].cells:
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True

doc.add_paragraph()

# === Section C: Ringkasan ===
doc.add_heading('C. Ringkasan Total', level=1)

table_c = doc.add_table(rows=4, cols=2)
table_c.style = 'Table Grid'
table_c.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_c = ['Kategori', 'Biaya/bulan']
for i, h in enumerate(headers_c):
    table_c.rows[0].cells[i].text = h
    table_c.rows[0].cells[i].paragraphs[0].runs[0].bold = True

data_c = [
    ['Infrastruktur', 'Rp 0'],
    ['Maintenance & Support', 'Rp 1.000.000'],
    ['TOTAL BIAYA BULANAN', 'Rp 1.000.000'],
]

for row_idx, row_data in enumerate(data_c):
    for col_idx, val in enumerate(row_data):
        table_c.rows[row_idx + 1].cells[col_idx].text = val

for cell in table_c.rows[3].cells:
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True

doc.add_paragraph()

# === Section D: Ketentuan ===
doc.add_heading('D. Ketentuan', level=1)

table_d = doc.add_table(rows=7, cols=3)
table_d.style = 'Table Grid'
table_d.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_d = ['No', 'Ketentuan', 'Detail']
for i, h in enumerate(headers_d):
    table_d.rows[0].cells[i].text = h
    table_d.rows[0].cells[i].paragraphs[0].runs[0].bold = True

data_d = [
    ['1', 'Kontrak minimum', '3 bulan'],
    ['2', 'Pembayaran', 'Di awal bulan'],
    ['3', 'Response time', 'Maks 1x24 jam (hari kerja)'],
    ['4', 'Bug fix di luar kuota', 'Rp 100.000/incident (>4 incident/bulan)'],
    ['5', 'Penambahan fitur baru', 'Biaya development terpisah sesuai scope'],
    ['6', 'Kapasitas', '40 pasien/hari, 1.200 visits/bulan'],
]

for row_idx, row_data in enumerate(data_d):
    for col_idx, val in enumerate(row_data):
        table_d.rows[row_idx + 1].cells[col_idx].text = val

doc.add_paragraph()

# === Section E: Opsional ===
doc.add_heading('E. Opsional (Upgrade di Masa Depan)', level=1)

table_e = doc.add_table(rows=4, cols=4)
table_e.style = 'Table Grid'
table_e.alignment = WD_TABLE_ALIGNMENT.CENTER

headers_e = ['No', 'Item', 'Biaya', 'Manfaat']
for i, h in enumerate(headers_e):
    table_e.rows[0].cells[i].text = h
    table_e.rows[0].cells[i].paragraphs[0].runs[0].bold = True

data_e = [
    ['1', 'Custom domain .id', 'Rp 150.000/tahun', 'URL lebih profesional'],
    ['2', 'WhatsApp Business API', 'Rp 500.000/bulan', 'Kirim hasil otomatis tanpa klik'],
    ['3', 'Dashboard Analytics', 'Rp 500.000 (one-time)', 'Visualisasi data pasien'],
]

for row_idx, row_data in enumerate(data_e):
    for col_idx, val in enumerate(row_data):
        table_e.rows[row_idx + 1].cells[col_idx].text = val

# Save
output_path = r'c:\KIRO AKUH\sinta\dashboard\questionnaire\Rekap_Biaya_Bulanan_Edukasi_Catin.docx'
doc.save(output_path)
print(f'Done: {output_path}')
