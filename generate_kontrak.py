"""Generate Surat Perjanjian Kerja Sama - Maintenance Web Edukasi Catin"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Header
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('SURAT PERJANJIAN KERJA SAMA')
run.bold = True
run.font.size = Pt(14)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('MAINTENANCE DAN SUPPORT SISTEM EDUKASI DIGITAL\nCALON PENGANTIN (CATIN)')
run.bold = True
run.font.size = Pt(12)

nomor = doc.add_paragraph()
nomor.alignment = WD_ALIGN_PARAGRAPH.CENTER
nomor.add_run('Nomor: ......../PKS/PUSKES-CKG/......./2026').font.size = Pt(11)

doc.add_paragraph()

# Pembuka
doc.add_paragraph(
    'Pada hari ini, ............... tanggal ............... bulan ............... '
    'tahun Dua Ribu Dua Puluh Enam (2026), bertempat di Puskesmas Cakung, '
    'telah disepakati Perjanjian Kerja Sama antara:'
)

doc.add_paragraph()

# Pihak 1
p1 = doc.add_paragraph()
p1.add_run('PIHAK PERTAMA:').bold = True
doc.add_paragraph('Nama\t\t: ........................................')
doc.add_paragraph('Jabatan\t\t: Kepala Puskesmas Cakung')
doc.add_paragraph('Alamat\t\t: ........................................')
doc.add_paragraph(
    'Dalam hal ini bertindak untuk dan atas nama Puskesmas Cakung, '
    'selanjutnya disebut PIHAK PERTAMA.'
)

doc.add_paragraph()

# Pihak 2
p2 = doc.add_paragraph()
p2.add_run('PIHAK KEDUA:').bold = True
doc.add_paragraph('Nama\t\t: Wiliam Fredrick Sakalessy')
doc.add_paragraph('Jabatan\t\t: IT Developer / Konsultan')
doc.add_paragraph('Alamat\t\t: ........................................')
doc.add_paragraph('No. HP/WA\t: ........................................')
doc.add_paragraph(
    'Dalam hal ini bertindak sebagai penyedia jasa maintenance dan support '
    'sistem edukasi digital, selanjutnya disebut PIHAK KEDUA.'
)

doc.add_paragraph()
doc.add_paragraph(
    'Kedua belah pihak sepakat untuk mengadakan perjanjian kerja sama '
    'dengan ketentuan sebagai berikut:'
)

doc.add_paragraph()

# Pasal 1
h1 = doc.add_paragraph()
h1.add_run('PASAL 1 — RUANG LINGKUP PEKERJAAN').bold = True
doc.add_paragraph(
    'PIHAK KEDUA menyediakan jasa maintenance dan technical support untuk '
    'Sistem Edukasi Digital Calon Pengantin (Web Edukasi Catin) yang diakses melalui '
    'URL: https://iamwfs30.github.io/educatin-puskes-cakung/'
)
doc.add_paragraph()
doc.add_paragraph('Ruang lingkup pekerjaan meliputi:')
doc.add_paragraph('1. Bug Fix — Perbaikan error/bug yang ditemukan saat operasional (maks 4 incident/bulan)')
doc.add_paragraph('2. Update Konten — Perubahan/penambahan soal, update video, perubahan teks (maks 2x/bulan)')
doc.add_paragraph('3. Monitoring — Pengecekan ketersediaan website secara mingguan')
doc.add_paragraph('4. Backup Data — Export rekap data Google Sheets dan backup source code (bulanan)')
doc.add_paragraph('5. Technical Support — Bantuan teknis via WhatsApp pada jam kerja (Senin-Jumat, 08.00-17.00 WIB)')

doc.add_paragraph()

# Pasal 2
h2 = doc.add_paragraph()
h2.add_run('PASAL 2 — JANGKA WAKTU').bold = True
doc.add_paragraph(
    '1. Perjanjian ini berlaku selama 6 (enam) bulan terhitung sejak tanggal '
    'ditandatanganinya perjanjian ini.'
)
doc.add_paragraph(
    '2. Perjanjian dapat diperpanjang dengan kesepakatan kedua belah pihak, '
    'yang disampaikan secara tertulis paling lambat 30 (tiga puluh) hari '
    'sebelum masa perjanjian berakhir.'
)

doc.add_paragraph()

# Pasal 3
h3 = doc.add_paragraph()
h3.add_run('PASAL 3 — BIAYA DAN PEMBAYARAN').bold = True
doc.add_paragraph('1. Nilai kontrak selama 6 bulan adalah sebesar:')
doc.add_paragraph('\tRp 1.000.000,- (Satu Juta Rupiah) per bulan')
doc.add_paragraph('\tTotal: Rp 6.000.000,- (Enam Juta Rupiah)')
doc.add_paragraph()
doc.add_paragraph('2. Pembayaran dilakukan dalam 2 (dua) termin:')
doc.add_paragraph('\ta. Termin 1: Rp 3.000.000,- (dibayarkan di awal kontrak)')
doc.add_paragraph('\tb. Termin 2: Rp 3.000.000,- (dibayarkan pada bulan ke-4)')
doc.add_paragraph()
doc.add_paragraph('3. Pembayaran dilakukan melalui transfer bank ke rekening PIHAK KEDUA:')
doc.add_paragraph('\tBank\t\t: ........................................')
doc.add_paragraph('\tNo. Rekening\t: ........................................')
doc.add_paragraph('\tAtas Nama\t: ........................................')

doc.add_paragraph()

# Pasal 4
h4 = doc.add_paragraph()
h4.add_run('PASAL 4 — SERVICE LEVEL AGREEMENT (SLA)').bold = True
doc.add_paragraph('1. Response time: Maksimal 1x24 jam pada hari kerja')
doc.add_paragraph('2. Bug fix critical (website down): Maksimal 4 jam')
doc.add_paragraph('3. Bug fix non-critical: Maksimal 2x24 jam hari kerja')
doc.add_paragraph('4. Update konten: Maksimal 3x24 jam hari kerja')
doc.add_paragraph('5. Uptime target: 99% per bulan')

doc.add_paragraph()

# Pasal 5
h5 = doc.add_paragraph()
h5.add_run('PASAL 5 — BIAYA TAMBAHAN').bold = True
doc.add_paragraph('1. Bug fix di luar kuota (lebih dari 4 incident/bulan): Rp 100.000,-/incident')
doc.add_paragraph('2. Penambahan fitur baru di luar scope maintenance: Dikenakan biaya development terpisah yang disepakati kedua belah pihak')
doc.add_paragraph('3. Update konten di luar kuota (lebih dari 2x/bulan): Rp 100.000,-/update')

doc.add_paragraph()

# Pasal 6
h6 = doc.add_paragraph()
h6.add_run('PASAL 6 — HAK DAN KEWAJIBAN').bold = True
doc.add_paragraph()
doc.add_paragraph('PIHAK PERTAMA berkewajiban:')
doc.add_paragraph('1. Melakukan pembayaran sesuai ketentuan Pasal 3')
doc.add_paragraph('2. Menyediakan informasi/konten yang diperlukan untuk update')
doc.add_paragraph('3. Melaporkan bug/masalah melalui WhatsApp')
doc.add_paragraph()
doc.add_paragraph('PIHAK KEDUA berkewajiban:')
doc.add_paragraph('1. Melaksanakan pekerjaan sesuai ruang lingkup Pasal 1')
doc.add_paragraph('2. Memenuhi SLA sesuai Pasal 4')
doc.add_paragraph('3. Menjaga kerahasiaan data pasien')
doc.add_paragraph('4. Memberikan laporan bulanan terkait status website')

doc.add_paragraph()

# Pasal 7
h7 = doc.add_paragraph()
h7.add_run('PASAL 7 — KERAHASIAAN DATA').bold = True
doc.add_paragraph(
    'PIHAK KEDUA wajib menjaga kerahasiaan seluruh data pasien yang tersimpan '
    'dalam sistem dan tidak diperkenankan membagikan, menjual, atau menggunakan '
    'data tersebut untuk kepentingan di luar perjanjian ini.'
)

doc.add_paragraph()

# Pasal 8
h8 = doc.add_paragraph()
h8.add_run('PASAL 8 — PENGAKHIRAN PERJANJIAN').bold = True
doc.add_paragraph('1. Perjanjian dapat diakhiri oleh salah satu pihak dengan pemberitahuan tertulis 30 hari sebelumnya.')
doc.add_paragraph('2. Dalam hal pengakhiran sebelum masa kontrak berakhir, sisa bulan yang belum berjalan akan di-refund secara proporsional.')
doc.add_paragraph('3. PIHAK KEDUA wajib menyerahkan seluruh akses, source code, dan data kepada PIHAK PERTAMA pada saat pengakhiran.')

doc.add_paragraph()

# Pasal 9
h9 = doc.add_paragraph()
h9.add_run('PASAL 9 — FORCE MAJEURE').bold = True
doc.add_paragraph(
    'Kedua belah pihak dibebaskan dari tanggung jawab atas keterlambatan atau '
    'kegagalan pelaksanaan kewajiban yang disebabkan oleh keadaan di luar '
    'kendali (bencana alam, kebijakan pemerintah, gangguan internet massal, dll).'
)

doc.add_paragraph()

# Pasal 10
h10 = doc.add_paragraph()
h10.add_run('PASAL 10 — PENYELESAIAN PERSELISIHAN').bold = True
doc.add_paragraph(
    'Apabila terjadi perselisihan, kedua belah pihak sepakat untuk menyelesaikan '
    'secara musyawarah mufakat. Apabila tidak tercapai kesepakatan, akan diselesaikan '
    'melalui jalur hukum yang berlaku di wilayah hukum Jakarta Timur.'
)

doc.add_paragraph()
doc.add_paragraph()

# Penutup
doc.add_paragraph(
    'Demikian perjanjian ini dibuat dalam rangkap 2 (dua) bermaterai cukup, '
    'masing-masing mempunyai kekuatan hukum yang sama.'
)

doc.add_paragraph()
doc.add_paragraph()

# Tanda tangan
table_ttd = doc.add_table(rows=5, cols=2)
table_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER

table_ttd.rows[0].cells[0].text = 'PIHAK PERTAMA,'
table_ttd.rows[0].cells[1].text = 'PIHAK KEDUA,'

table_ttd.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table_ttd.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

table_ttd.rows[3].cells[0].text = '(........................................)'
table_ttd.rows[3].cells[1].text = '(Wiliam Fredrick Sakalessy)'

table_ttd.rows[3].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table_ttd.rows[3].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

table_ttd.rows[4].cells[0].text = 'Kepala Puskesmas Cakung'
table_ttd.rows[4].cells[1].text = 'IT Developer / Konsultan'

table_ttd.rows[4].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table_ttd.rows[4].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Save
output_path = r'c:\KIRO AKUH\sinta\dashboard\questionnaire\Kontrak_Maintenance_Edukasi_Catin.docx'
doc.save(output_path)
print(f'Done: {output_path}')
